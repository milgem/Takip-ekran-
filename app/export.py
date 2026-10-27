"""Analiz tablolarını Excel (.xlsx) dosyası olarak dışa aktarır."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="1F4D3F")
RED = PatternFill("solid", fgColor="FDECEA")
AMBER = PatternFill("solid", fgColor="FDF3E2")
TOTAL = PatternFill("solid", fgColor="F0EFEA")
THIN = Side(style="thin", color="D9D8D2")


def _font(**k):
    return Font(name="Arial", **k)


def _sheet(wb, title, headers, rows, widths=None, money=(), ints=(), dec=(), first=False, note=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = _font(italic=True, size=9, color="666666")
        r0 = 3
    for i, h in enumerate(headers, 1):
        c = ws.cell(r0, i, h)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for ri, row in enumerate(rows, r0 + 1):
        fill = row.pop("_fill", None) if isinstance(row, dict) else None
        vals = row["v"] if isinstance(row, dict) else row
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v)
            c.font = _font(bold=bool(isinstance(row, dict) and row.get("bold")))
            c.border = Border(bottom=THIN)
            if ci - 1 in money:
                c.number_format = '#,##0 "₺";-#,##0 "₺";-'
            elif ci - 1 in ints:
                c.number_format = '#,##0;-#,##0;-'
            elif ci - 1 in dec:
                c.number_format = '#,##0.0;-#,##0.0;-'
            if fill:
                c.fill = fill
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + max(len(rows), 1)}"
    return ws


def _days(v):
    return None if v is None else round(max(v, 0), 1)


def _bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def assort_xlsx(firm_name: str, d: dict, only: list | None = None) -> bytes:
    """Asorti önerisi: özet, sipariş formu (model × beden) ve beden detayı."""
    from .deep import _size_key
    groups = only if only is not None else d["groups"]
    wb = Workbook()
    note = (f"{firm_name} · sipariş / üretim önerisi · {datetime.now():%d.%m.%Y %H:%M} · son {d['window']} günün satışına göre, "
            f"tedarik süresi {d['lead']} gün, istenen stok {d['cover']} gün")
    _sheet(wb, "Özet", ["Model", "Renk", "Asorti (beden:adet)", "Seri", "1 seri (adet)", "Verilecek toplam", "30 gün satış",
                         "En erken biten (gün)", "Acil"],
           [{"v": [g["product"], g["color"], g["ratio_text"], g["series"], g["series_size"], g["total"], g["sold"],
                   _days(g["min_days"]), "Evet" if g["urgent"] else ""], "_fill": RED if g["urgent"] else None} for g in groups],
           widths=[52, 14, 46, 8, 10, 14, 12, 14, 8], ints=(3, 4, 5, 6), dec=(7,), first=True, note=note)
    # Sipariş formu: model × beden. Harf bedenler (S, M, L…) ve numara bedenler (46, 48…) ayrı sayfalarda,
    # böylece her sayfada sadece o ürünlerin kullandığı beden sütunları olur.
    fam_name = {0: "numara", 1: "harf", 2: "renk-tek beden"}
    fams = {}
    for g in groups:
        made = [r for r in g["rows"] if r["make"]]
        if not made:
            continue
        cats = [_size_key(r["size"])[0] for r in made]
        fams.setdefault(max(set(cats), key=cats.count), []).append(g)
    for fam, gs in sorted(fams.items()):
        sizes = sorted({r["size"] for g in gs for r in g["rows"] if r["make"]}, key=_size_key)
        form = []
        for g in gs:
            by = {r["size"]: r["make"] for r in g["rows"]}
            form.append([g["product"], g["color"]] + [by.get(sz) or None for sz in sizes] + [g["total"]])
        tot = ["TOPLAM", ""] + [sum((f[2 + i] or 0) for f in form) for i in range(len(sizes))] + [sum(f[-1] for f in form)]
        form.append({"v": tot, "bold": True, "_fill": TOTAL})
        title = "Sipariş formu" if len(fams) == 1 else f"Sipariş formu ({fam_name[fam]})"
        _sheet(wb, title, ["Model", "Renk"] + sizes + ["Toplam"], form,
               widths=[52, 14] + [7] * len(sizes) + [10], ints=tuple(range(2, len(sizes) + 3)), note=note)
    det = []
    for g in groups:
        for r in g["rows"]:
            det.append({"v": [g["product"], g["color"], r["size"], r["sku"], r["sold"], round(r["daily"], 2), r["stock"],
                              _days(r["days_left"]), round(r["need"]), r["ratio"], r["make"], _days(r["days_after"])],
                        "_fill": RED if r["stock"] <= 0 else None})
    _sheet(wb, "Beden detayı", ["Model", "Renk", "Beden", "Stok kodu", "30 gün satış", "Günlük", "Stok", "Kalan gün", "İhtiyaç",
                                "Seride", "Verilecek", "Sonra kaç gün yeter"], det,
           widths=[52, 14, 8, 18, 11, 9, 9, 10, 9, 8, 10, 12], ints=(4, 6, 8, 9, 10), dec=(5, 7, 11), note=note)
    return _bytes(wb)


def sizes_xlsx(firm_name: str, d: dict) -> bytes:
    wb = Workbook()
    note = f"{firm_name} · beden bazında sipariş önerisi · {datetime.now():%d.%m.%Y %H:%M}"
    label = {"acil": "Acil", "yakin": "Yakında", "plan": "Planla"}
    _sheet(wb, "Sipariş önerisi", ["Marka", "Ürün", "Beden", "Stok kodu", f"{d['window']} gün satış", "Günlük", "Stok", "Kalan gün",
                                   "Önerilen", "Değer", "Durum"],
           [{"v": [r["brand"], r["product"], r["variant"], r["sku"], r["sold"], round(r["daily"], 2), r["stock"], _days(r["days_left"]),
                   r["suggest"], round(r["value"]), label[r["urgency"]]], "_fill": RED if r["urgency"] == "acil" else None}
            for r in d["rows"]],
           widths=[20, 52, 14, 18, 11, 9, 9, 10, 10, 12, 10], ints=(4, 6, 8), dec=(5, 7), money=(9,), first=True, note=note)
    return _bytes(wb)


def dead_xlsx(firm_name: str, d: dict) -> bytes:
    wb = Workbook()
    note = f"{firm_name} · son {d['days']} günde satmayan stok · {datetime.now():%d.%m.%Y %H:%M}"
    _sheet(wb, "Ürünler", ["Ürün", "Marka", "Kategori", "Beden sayısı", "Stok", "Değer (satış fiyatı)", "Son satış"],
           [[r["product"], r["brand"], r["category"], r["variant_count"], r["stock"], round(r["value"]),
             r["last"].strftime("%d.%m.%Y") if r["last"] else ""] for r in d["rows"]],
           widths=[52, 18, 22, 11, 10, 16, 12], ints=(3, 4), money=(5,), first=True, note=note)
    _sheet(wb, "Bedenler", ["Ürün", "Marka", "Beden", "Stok kodu", "Stok", "Fiyat", "Son satış"],
           [[r["product"], r["brand"], x["variant"], x["sku"], x["stock"], x["price"], x["last"].strftime("%d.%m.%Y") if x["last"] else ""]
            for r in d["rows"] for x in r["variants"]],
           widths=[52, 18, 14, 18, 10, 12, 12], ints=(4,), money=(5,), note=note)
    return _bytes(wb)


def profit_xlsx(firm_name: str, d: dict, start, end) -> bytes:
    wb = Workbook()
    note = f"{firm_name} · kanal kârlılığı · {start:%d.%m.%Y} – {end:%d.%m.%Y}"
    _sheet(wb, "Kanallar", ["Kanal", "Sipariş", "Ciro", "Komisyon", "Maliyet", "Reklam", "Net kâr", "Net marj %", "Maliyet verisi %"],
           [[r["channel"], r["orders"], round(r["revenue"]), round(r["commission"]),
             round(r["cost"]) if r["cost"] is not None else None, round(r["ad"]),
             round(r["net"]) if r["net"] is not None else None,
             round(r["net_margin"], 1) if r["net_margin"] is not None else None, round(r["coverage"])] for r in d["rows"]],
           widths=[18, 10, 14, 14, 14, 12, 14, 11, 13], ints=(1,), money=(2, 3, 4, 5, 6), dec=(7,), first=True, note=note)
    _sheet(wb, "Ürünler", ["Ürün", "Adet", "Ciro", "Maliyet", "Komisyon", "Brüt kâr", "Marj %"],
           [[p["product"], p["qty"], round(p["revenue"]), round(p["cost"]), round(p["commission"]), round(p["profit"]),
             round(p["margin"], 1) if p["margin"] is not None else None] for p in d["products"]],
           widths=[52, 9, 14, 14, 12, 14, 9], ints=(1,), money=(2, 3, 4, 5), dec=(6,), note=note)
    return _bytes(wb)
